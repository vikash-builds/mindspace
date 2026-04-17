from pptx import Presentation

def parse_pptx(file_path):
    text = ""
    try:
        prs = Presentation(file_path)
        for i, slide in enumerate(prs.slides):
            text += f"--- Slide {i+1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
            text += "\n"
    except Exception as e:
        print(f"Error parsing PPTX: {e}")
    return text
